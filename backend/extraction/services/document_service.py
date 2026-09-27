"""
ModelLens — Document Extraction Service

Extracts text from uploaded files (PDF, DOCX, TXT, MD, CSV, JSON, and generic text),
splits the text into overlapping word-based chunks, and returns structured Chunk objects.
"""

from __future__ import annotations

import logging
from pathlib import Path

import pymupdf as fitz  # PyMuPDF

try:
    import docx  # python-docx
except ImportError:
    docx = None  # type: ignore

from extraction.config import CHUNK_SIZE_WORDS, CHUNK_OVERLAP_WORDS
from extraction.schemas import Chunk

logger = logging.getLogger(__name__)

# Typical page size in words for documents without native physical pages (e.g. txt, docx, md)
WORDS_PER_PAGE = 500


def _extract_pdf(file_path: Path) -> list[dict]:
    """Extract page-by-page text from a PDF file using PyMuPDF."""
    try:
        doc = fitz.open(str(file_path))
    except Exception as exc:
        raise ValueError(f"Cannot open PDF: {exc}") from exc

    pages: list[dict] = []
    for page_number, page in enumerate(doc, start=1):
        text = page.get_text().strip()
        if text:
            pages.append({"page": page_number, "text": text})

    doc.close()
    return pages


def _extract_docx(file_path: Path) -> list[dict]:
    """Extract text from a Word (.docx) file and paginate by word count."""
    if docx is None:
        raise ValueError("Word (.docx) support requires 'python-docx'. Please run 'pip install python-docx'.")

    try:
        doc = docx.Document(str(file_path))
    except Exception as exc:
        raise ValueError(f"Cannot open Word document: {exc}") from exc

    blocks: list[str] = []
    for p in doc.paragraphs:
        t = p.text.strip()
        if t:
            blocks.append(t)

    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                blocks.append(row_text)

    if not blocks:
        return []

    return _paginate_text("\n\n".join(blocks))


def _extract_text_file(file_path: Path) -> list[dict]:
    """Read a plain text, markdown, csv, json, or other text-based file."""
    raw_bytes = file_path.read_bytes()

    # Quick check for binary files (e.g. images, executables, compiled files)
    if b"\x00" in raw_bytes[:1024]:
        raise ValueError(f"File '{file_path.name}' is binary or contains no extractable text.")

    text: str | None = None
    for enc in ("utf-8", "utf-8-sig", "latin-1"):
        try:
            text = raw_bytes.decode(enc)
            break
        except UnicodeDecodeError:
            continue

    if not text:
        raise ValueError(f"Unable to decode text from '{file_path.name}'.")

    text = text.strip()
    if not text:
        return []

    return _paginate_text(text)


def _paginate_text(full_text: str) -> list[dict]:
    """Split continuous text into ~WORDS_PER_PAGE blocks to preserve page numbers."""
    words = full_text.split()
    if not words:
        return []

    pages: list[dict] = []
    page_idx = 1
    start = 0

    while start < len(words):
        end = start + WORDS_PER_PAGE
        page_words = words[start:end]
        pages.append({"page": page_idx, "text": " ".join(page_words)})
        page_idx += 1
        start = end

    return pages


# ── Public API ──────────────────────────────────────────────────────────────


def extract_text_by_page(file_path: str | Path) -> list[dict]:
    """
    Open any document file and return a list of dicts with page-level text.

    Supports:
      - PDF (.pdf)
      - Word (.docx)
      - Plain text (.txt, .md, .csv, .json, .py, .yaml, .xml, .html, logs, etc.)
      - Any generic decodable text file

    Returns
    -------
    [
        {"page": 1, "text": "..."},
        {"page": 2, "text": "..."},
        ...
    ]

    Raises
    ------
    ValueError  – if the file cannot be opened or contains no extractable text.
    """
    file_path = Path(file_path)
    if not file_path.exists():
        raise ValueError(f"File not found: {file_path}")

    suffix = file_path.suffix.lower()

    # Check for PDF (by extension or header magic bytes)
    try:
        with open(file_path, "rb") as f:
            header = f.read(5)
    except OSError as exc:
        raise ValueError(f"Cannot read file: {exc}") from exc

    pages: list[dict] = []

    if suffix == ".pdf" or header.startswith(b"%PDF"):
        pages = _extract_pdf(file_path)
    elif suffix == ".docx":
        pages = _extract_docx(file_path)
    else:
        # Try as text / markdown / code / json / etc.
        try:
            pages = _extract_text_file(file_path)
        except Exception:
            # If text extraction fails, attempt fitz as fallback
            try:
                pages = _extract_pdf(file_path)
            except Exception:
                raise ValueError(f"File '{file_path.name}' contains no extractable text.")

    if not pages:
        raise ValueError(f"File '{file_path.name}' contains no extractable text.")

    logger.info("Extracted %d pages from %s", len(pages), file_path.name)
    return pages


def build_chunks(pages: list[dict]) -> list[Chunk]:
    """
    Split page texts into overlapping word-based chunks (~500-800 words).

    Each chunk retains the page number it primarily belongs to.
    Overlap is applied so that information at boundaries is not lost.
    """
    chunks: list[Chunk] = []
    chunk_counter = 0

    for page_info in pages:
        page_num = page_info["page"]
        words = page_info["text"].split()

        if not words:
            continue

        start = 0
        while start < len(words):
            end = start + CHUNK_SIZE_WORDS
            chunk_words = words[start:end]
            chunk_text = " ".join(chunk_words)

            chunk_counter += 1
            chunks.append(
                Chunk(
                    chunk_id=f"chunk_{chunk_counter:03d}",
                    text=chunk_text,
                    page=page_num,
                )
            )

            # Advance by (chunk_size - overlap) so consecutive chunks overlap
            step = max(CHUNK_SIZE_WORDS - CHUNK_OVERLAP_WORDS, 1)
            start += step

    logger.info("Built %d chunks from %d pages", len(chunks), len(pages))
    return chunks


def extract_and_chunk(file_path: str | Path) -> tuple[int, list[Chunk]]:
    """
    Convenience wrapper: extract text then chunk.

    Returns
    -------
    (page_count, chunks)
    """
    pages = extract_text_by_page(file_path)
    page_count = len(pages)
    chunks = build_chunks(pages)
    return page_count, chunks
